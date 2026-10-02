from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUTPUT = "/Users/nikos/Desktop/peireas/D1.1_citation_audit.docx"

issues = [
    ("Missing reference", "Bergé (2018), Section 4.4.4", "No corresponding entry appears in the reference list. Add the complete fixest citation."),
    ("Year mismatch", "Oser (2021)", "The bibliography lists Oser (2022). Verify and use the correct year consistently."),
    ("Author mismatch", "Alscher and Costa (2025)", "The bibliography lists Alscher and Jana (2025). Verify the authors against the original source."),
    ("Author name mismatch", "Hameleers and Ortiz (2024)", "The bibliography lists Hameleers and Garnier Ortiz (2024). Use the same surname form in both places."),
    ("Reversed author order", "Lefkofridi and Giger (2014)", "The bibliography lists Giger and Lefkofridi (2014). Author order must agree with the publication."),
    ("Reversed author order", "Potter and Wetherell (1996)", "The bibliography lists Wetherell and Potter (1996). Author order must agree with the publication."),
    ("Author name inconsistency", "Borge Bravo et al. (2019)", "The bibliography lists Borge, Balcells, and Padró-Solanet (2019). Verify whether the first author's surname is Borge or Borge Bravo."),
    ("Author name inconsistency", "Coronel, Colón Amill, and Drouin (2019)", "Elsewhere the source appears as Coronel, Amill, and Drouin, and the bibliography uses Amill. Standardise the second author's surname."),
    ("Ambiguous same year", "Ringe (2022)", "Two Ringe publications from 2022 are listed. Assign 2022a and 2022b and apply the suffixes to every relevant citation."),
    ("Ambiguous same year", "Verhasselt (2025)", "Two Verhasselt publications from 2025 are listed. Assign 2025a and 2025b and apply the suffixes consistently."),
    ("Insufficient identification", "The 2024 EU survey on Europeans and their languages", "A European Commission (2024) entry exists, but the paragraph should cite it explicitly as (European Commission, 2024)."),
]

minor_points = [
    "(Fricker 2007) is missing a comma between the author and year.",
    "Lai & beh (2025) should capitalise Beh.",
    "Benjamini-Hochberg (1995) should normally be written as Benjamini and Hochberg (1995) when identifying the authors.",
    "Names such as Marková and Hyvärinen lose their diacritics in some in-text citations.",
    "García and Li Wei or Wei is not presented consistently across the text and reference list.",
]

incomplete = [
    "Borge, Balcells, and Padró-Solanet (2019)",
    "Cooper (2023)",
    "Erdocia (2023)",
    "Finocchiaro and Godden (2011)",
    "Mead (2018)",
    "Rocci (2006)",
    "Van der Velden (2018)",
    "Kelly, Tilley, and Oskarsson (2025)",
]

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)

def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)

def set_font(run, name="Aptos", size=10.5, bold=False, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.08
    set_font(p.add_run(text), size=10.5)
    return p

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.12

for style_name, size in (("Title", 23), ("Heading 1", 15), ("Heading 2", 12)):
    style = styles[style_name]
    style.font.name = "Aptos Display" if style_name != "Heading 2" else "Aptos"
    style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)

styles["Title"].paragraph_format.space_after = Pt(8)
styles["Heading 1"].paragraph_format.space_before = Pt(15)
styles["Heading 1"].paragraph_format.space_after = Pt(7)
styles["Heading 2"].paragraph_format.space_before = Pt(11)
styles["Heading 2"].paragraph_format.space_after = Pt(5)
title_ppr = styles["Title"]._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

title = doc.add_paragraph(style="Title")
set_font(title.add_run("Citation Audit for D1 1 Tracked Document"), name="Aptos Display", size=23, bold=True)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(14)
set_font(subtitle.add_run("Review of in-text citations against the reference list"), size=11.5, color="536273")

intro = doc.add_paragraph()
intro.add_run("Conclusion  ").bold = True
intro.add_run("Not every citation is cleanly or unambiguously identified in the document. Most literature-review citations have a recognisable reference-list entry, but several missing, conflicting, ambiguous, or incomplete records should be corrected before finalisation.")

scope = doc.add_paragraph()
scope.add_run("Scope  ").bold = True
scope.add_run("The check used the current text in D1.1_tracked.docx, including tracked insertions and excluding tracked deletions.")

doc.add_heading("Main citation and reference issues", level=1)
table = doc.add_table(rows=1, cols=3)
table.autofit = False
table.columns[0].width = Inches(1.35)
table.columns[1].width = Inches(2.25)
table.columns[2].width = Inches(3.36)
set_table_borders(table)
headers = ["Issue", "In-text citation", "Reference-list finding and action"]
for cell, text in zip(table.rows[0].cells, headers):
    set_cell_shading(cell, "1F4E78")
    set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    set_font(p.add_run(text), size=9.3, bold=True, color="FFFFFF")
set_repeat_table_header(table.rows[0])

for idx, row_data in enumerate(issues):
    row = table.add_row()
    for col, (cell, text) in enumerate(zip(row.cells, row_data)):
        cell.width = table.columns[col].width
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        if idx % 2 == 1:
            set_cell_shading(cell, "EEF4F8")
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.04
        set_font(p.add_run(text), size=9.2, bold=(col == 0))

doc.add_page_break()
doc.add_heading("Smaller citation-format problems", level=1)
for point in minor_points:
    add_bullet(doc, point)

doc.add_heading("Bibliography entries needing fuller identification", level=1)
p = doc.add_paragraph("The following cited works appear in the reference list but do not contain enough publication information to be located confidently:")
p.paragraph_format.space_after = Pt(5)
for item in incomplete:
    add_bullet(doc, item)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(5)
p.add_run("Required completion  ").bold = True
p.add_run("Add the journal, book, publisher, report institution, DOI, URL, or other publication information appropriate to each source.")

doc.add_heading("Overall assessment", level=1)
p = doc.add_paragraph()
p.add_run("The bibliography is mostly connected to the literature review, but the document is not yet citation-clean. ").bold = True
p.add_run("The clearest substantive omission is Bergé (2018). The most important ambiguities are the two Ringe (2022) and two Verhasselt (2025) entries. The author and year conflicts listed above should also be resolved against the original publications before submission.")

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(fp.add_run("D1.1 citation audit"), size=8.5, color="6B7280")

doc.save(OUTPUT)
print(OUTPUT)
